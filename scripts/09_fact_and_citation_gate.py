#!/usr/bin/env python3
"""
Stage 7B: Pre-Delivery Fact, Peptide-Tier, Overclaim, & Citation Gate
=====================================================================
The mandatory final verification gate before any report (.md / .txt), schematic
(.svg), or presentation (.pptx) is delivered to the user.

Automated verification checks against a canonical `facts.json` (or a results CSV):
  1. Gene-Number Reconciliation: every signed number written next to a gene
     (``GENE +1.23``, ``GENE (−0.87``, ``GENE log2FC = -2.10``) must match one of
     that gene's canonical metrics within ``tolerance`` (default 0.02).
  1b. Quarantined-value leak: values listed in ``quarantined_outlier_values`` may
     only be cited together with an exclusion disclosure ("outlier", "excluded",
     "quarantined", or the name of a sample listed in ``facts["quarantined_samples"]``).
  1c. Quadrant / category contradiction against ``categories``.
  2. Single-Peptide Promotion Audit: genes with ``unique_peptides == 1`` need an explicit
     1-peptide / exploratory disclosure on the same line. Missing peptide counts are
     treated as UNKNOWN (warning), never silently defaulted to 2.
  3. Overclaim Linter: causal verbs ("proves", "confirms", "demonstrates", ...) on the
     same line as a gene graded ``E3`` or a contrast whose ``empirical_null_fdr > 0.25``.
  4. Live Bibliographic PMID/DOI Check via the EuropePMC REST API (verified TLS,
     8 s timeout). Offline -> ``UNVERIFIED`` (status WARN) unless ``allow_offline``.

facts.json schema (all keys optional except ``genes``)::

    {
      "genes": {
        "GENE": {"unique_peptides": 3 | null,
                 "metrics": {"log2FC": 1.23, ...},
                 "categories": {"quadrant": "Q1_concordant_up"},
                 "evidence_grade": "E1" | "E2" | "E3",
                 "contrast": "KD_vs_CTRL",
                 "quarantined_outlier_values": [19.88]}
      },
      "contrasts": {"KD_vs_CTRL": {"empirical_null_fdr": 0.04, "evidence_grade": "E1"}},
      "quarantined_samples": ["CTRL_Rep2"],
      "single_peptide_exempt_contexts": ["regulon"],
      "ambiguous_tokens": ["S1", "C1"]
    }

Exit codes (CLI): 0 = PASS / WARN, 1 = FAIL.
"""

from __future__ import annotations

import argparse
import json
import re
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

import pandas as pd

EUROPEPMC_SEARCH = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
NUMBER_TOLERANCE = 0.02
E3_FDR_THRESHOLD = 0.25

CAUSAL_VERBS = (
    "proves", "proved", "proven", "prove",
    "confirms", "confirmed", "confirm",
    "demonstrates", "demonstrated", "demonstrate",
    "establishes", "established",
    "证明", "证实", "确证",
)

GENERIC_EXCLUSION_TOKENS = ("outlier", "excluded", "quarantined", "quarantine", "剔除", "异常")

SINGLE_PEP_DISCLOSURE_TOKENS = (
    "1-pep", "1 pep", "single-peptide", "single peptide",
    "unique_peptides = 1", "unique_peptides == 1", "unique_peptides=1",
    "| 1 |", "exploratory", "单肽段", "1个独特肽段",
)

QUADRANT_TOKENS = (
    "Q1_concordant_up", "Q2_hyper_secreted", "Q3_concordant_down",
    "Q4_intracellular_retained", "Q4_tissue_retained", "Q_fluid_depleted",
)

# Generic workflow labels that collide with short gene-like tokens (not lab specific).
DEFAULT_AMBIGUOUS_TOKENS = frozenset(
    {"TF", "C1", "C2", "C3", "C4", "C5", "C6", "S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8", "S9",
     "A1", "A2", "A3", "A4", "E1", "E2", "E3", "M1", "M2", "M3", "M4", "M5", "M6"}
)
# Generic contexts in which a gene name denotes a regulator/regulon, not a measured protein.
DEFAULT_SINGLE_PEP_EXEMPT_CONTEXTS = ("regulon", "transcription factor", "upstream regulator")

_SIGNED_NUM = re.compile(r"(?<![\w.])([+\-−±])\s?(\d+\.\d{2,})(?![\d])")


# ---------------------------------------------------------------------------
# Facts construction
# ---------------------------------------------------------------------------
def build_facts_dictionary_from_csv(
    csv_path: str | Path,
    gene_col: str = "gene_symbol",
    pep_col: str = "unique_peptides",
    numeric_cols: Optional[Sequence[str]] = None,
    category_cols: Optional[Sequence[str]] = None,
    quarantined_values_by_gene: Optional[Dict[str, List[float]]] = None,
    quarantined_samples: Optional[Sequence[str]] = None,
) -> Dict[str, Any]:
    """Compile a canonical `facts.json` dictionary from a processed results CSV.

    Missing / non-positive peptide counts are stored as ``None`` (unknown).
    """
    df = pd.read_csv(csv_path)
    if gene_col not in df.columns:
        gene_col = next((c for c in df.columns if "gene" in c.lower()), df.columns[0])
    if pep_col not in df.columns:
        pep_col = next(
            (c for c in df.columns if "unique_pep" in c.lower()),
            next((c for c in df.columns if "pep" in c.lower() and "tier" not in c.lower()), ""),
        )

    if numeric_cols is None:
        numeric_cols = [c for c in df.columns if c != pep_col and pd.api.types.is_numeric_dtype(df[c])]
    if category_cols is None:
        category_cols = [
            c for c in df.columns
            if any(k in c.lower() for k in ["quadrant", "sector", "tier", "class", "regime"])
        ]
    grade_col = next((c for c in df.columns if c.lower() in ("evidence_grade", "grade")), None)

    genes_dict: Dict[str, Any] = {}
    for _, row in df.iterrows():
        g = str(row.get(gene_col, "")).strip()
        if not g or g.lower() == "nan":
            continue
        peps: Optional[int] = None
        if pep_col:
            raw_pep = pd.to_numeric(row.get(pep_col), errors="coerce")
            if pd.notna(raw_pep) and raw_pep > 0:
                peps = int(raw_pep)
        num_map = {}
        for nc in numeric_cols:
            val = pd.to_numeric(row.get(nc), errors="coerce")
            if pd.notna(val):
                num_map[nc] = round(float(val), 4)
        cat_map = {cc: str(row.get(cc, "")) for cc in category_cols if pd.notna(row.get(cc))}
        entry: Dict[str, Any] = {
            "unique_peptides": peps,
            "metrics": num_map,
            "categories": cat_map,
            "quarantined_outlier_values": (quarantined_values_by_gene or {}).get(g, []),
        }
        if grade_col and pd.notna(row.get(grade_col)):
            entry["evidence_grade"] = str(row.get(grade_col))
        genes_dict[g] = entry

    return {
        "source_csv": str(csv_path),
        "genes": genes_dict,
        "quarantined_samples": list(quarantined_samples or []),
    }


# ---------------------------------------------------------------------------
# Text extraction (.md / .txt / .svg / .pptx)
# ---------------------------------------------------------------------------
def _svg_text_lines(path: Path) -> List[str]:
    try:
        root = ET.parse(str(path)).getroot()
    except ET.ParseError:
        return re.sub(r"<[^>]+>", "\n", path.read_text(encoding="utf-8", errors="ignore")).splitlines()
    lines: List[str] = []
    for el in root.iter():
        tag = el.tag.split("}")[-1]
        if tag == "text":
            txt = "".join(el.itertext()).strip()
            if txt:
                lines.append(txt)
    return lines


def _pptx_text_lines(path: Path) -> List[str]:
    try:
        from pptx import Presentation  # type: ignore
    except Exception as exc:  # pragma: no cover - dependency missing
        raise RuntimeError(f"python-pptx is required to audit {path.name}: {exc}") from exc
    prs = Presentation(str(path))
    lines: List[str] = []

    def _walk(shapes: Iterable[Any], slide_idx: int) -> None:
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):  # group
                _walk(sh.shapes, slide_idx)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    cells = [c.text.strip() for c in r.cells]
                    if any(cells):
                        lines.append("| " + " | ".join(cells) + " |")

    for i, slide in enumerate(prs.slides, start=1):
        _walk(slide.shapes, i)
        if slide.has_notes_slide:
            for para in slide.notes_slide.notes_text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    lines.append(t)
    return lines


def extract_text_lines(path: str | Path) -> List[str]:
    """Return auditable text lines for markdown/plain text, SVG, or PPTX files."""
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".svg":
        return _svg_text_lines(p)
    if suf == ".pptx":
        return _pptx_text_lines(p)
    return p.read_text(encoding="utf-8").splitlines()


# ---------------------------------------------------------------------------
# Citation verification
# ---------------------------------------------------------------------------
def _ssl_context() -> ssl.SSLContext:
    try:
        import certifi  # type: ignore

        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        return ssl.create_default_context()


def verify_citation_europepmc(kind: str, ident: str, timeout: float = 8.0) -> Dict[str, Any]:
    """Verify a PMID or DOI via EuropePMC. Returns status VERIFIED / NOT_FOUND / UNVERIFIED."""
    if kind == "pmid":
        query = f"EXT_ID:{ident} AND SRC:MED"
    else:
        query = f'DOI:"{ident}"'
    url = f"{EUROPEPMC_SEARCH}?{urllib.parse.urlencode({'query': query, 'format': 'json', 'pageSize': 1})}"
    try:
        with urllib.request.urlopen(url, timeout=timeout, context=_ssl_context()) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
        return {"kind": kind, "id": ident, "status": "UNVERIFIED", "reason": f"network/parse error: {exc}"}
    hits = int(payload.get("hitCount", 0) or 0)
    if hits > 0:
        rec = (payload.get("resultList", {}).get("result") or [{}])[0]
        return {"kind": kind, "id": ident, "status": "VERIFIED", "title": rec.get("title", "")}
    return {"kind": kind, "id": ident, "status": "NOT_FOUND"}


# ---------------------------------------------------------------------------
# Core audit
# ---------------------------------------------------------------------------
def _gene_grade(ginfo: Dict[str, Any]) -> str:
    g = ginfo.get("evidence_grade") or ginfo.get("categories", {}).get("evidence_grade") or ""
    return str(g).upper()


def _to_float_signed(sign: str, digits: str) -> List[float]:
    v = float(digits)
    if sign in ("-", "−"):
        return [-v]
    if sign == "±":
        return [v, -v]
    return [v]


def audit_text_lines_against_facts(
    lines: Sequence[str],
    facts: Dict[str, Any],
    allow_untagged_single_peptide: bool = False,
    tolerance: float = NUMBER_TOLERANCE,
    verify_citations: bool = True,
    allow_offline: bool = False,
    source: str = "<text>",
) -> Dict[str, Any]:
    """Audit already-extracted text lines against a facts dictionary."""
    genes_db: Dict[str, Any] = facts.get("genes", {}) or {}
    contrasts_db: Dict[str, Any] = facts.get("contrasts", {}) or {}
    quarantined_samples = [str(s).lower() for s in facts.get("quarantined_samples", []) or []]
    exclusion_tokens = tuple(GENERIC_EXCLUSION_TOKENS) + tuple(quarantined_samples)
    ambiguous = set(DEFAULT_AMBIGUOUS_TOKENS) | set(facts.get("ambiguous_tokens", []) or [])
    exempt_ctx = tuple(DEFAULT_SINGLE_PEP_EXEMPT_CONTEXTS) + tuple(
        str(t).lower() for t in facts.get("single_peptide_exempt_contexts", []) or []
    )

    critical_errors: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    unknown_pep_warned: set = set()

    gene_patterns = {
        g: re.compile(rf"(?<![A-Za-z0-9_\-]){re.escape(g)}(?![A-Za-z0-9_])")
        for g in genes_db if len(g) >= 2
    }
    weak_contrasts = {
        name for name, c in contrasts_db.items()
        if (pd.notna(pd.to_numeric(c.get("empirical_null_fdr"), errors="coerce"))
            and float(c.get("empirical_null_fdr")) > E3_FDR_THRESHOLD)
        or str(c.get("evidence_grade", "")).upper() == "E3"
    }

    def _err(line_no: int, typ: str, gene: str, msg: str, line: str) -> None:
        critical_errors.append({"line": line_no, "type": typ, "gene": gene, "message": msg, "excerpt": line.strip()[:160]})

    def _warn(line_no: int, typ: str, gene: str, msg: str, line: str = "") -> None:
        warnings.append({"line": line_no, "type": typ, "gene": gene, "message": msg, "excerpt": line.strip()[:160]})

    for line_no, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        line_lower = line.lower()
        has_causal = any(re.search(rf"(?<!\w){re.escape(v)}(?!\w)", line_lower) for v in CAUSAL_VERBS)

        # Locate every gene mention (for nearest-preceding-gene number attribution)
        mentions: List[tuple] = []
        for gene, pat in gene_patterns.items():
            for m in pat.finditer(line):
                mentions.append((m.start(), m.end(), gene))
        mentions.sort()
        present = []
        for _, _, g in mentions:
            if g not in present:
                present.append(g)

        for gene in present:
            ginfo = genes_db[gene]

            # 1B: quarantined value leak
            for q_val in ginfo.get("quarantined_outlier_values", []) or []:
                q_str = f"{float(q_val):.2f}"
                if q_str in line and not any(tok in line_lower for tok in exclusion_tokens):
                    _err(line_no, "QUARANTINED_OUTLIER_VALUE_LEAK", gene,
                         f"Line {line_no} cites quarantined outlier value {q_str} for {gene} without noting it is from an excluded outlier sample.", line)

            # 1C: quadrant / category contradiction
            true_cats = set((ginfo.get("categories") or {}).values())
            for q_tok in QUADRANT_TOKENS:
                if q_tok in line and true_cats and q_tok not in true_cats:
                    if not any(neg in line_lower for neg in ["not ", "rather than", "而非", "修正"]):
                        _err(line_no, "CATEGORY_QUADRANT_MISMATCH", gene,
                             f"Line {line_no} associates {gene} with '{q_tok}', but canonical table classifies {gene} as {sorted(true_cats)}.", line)

            # 2: single-peptide promotion (unknown -> warning)
            raw_pep = ginfo.get("unique_peptides", None)
            peps = pd.to_numeric(raw_pep, errors="coerce") if raw_pep is not None else float("nan")
            if gene in ambiguous or any(t in line_lower for t in exempt_ctx):
                pass
            elif pd.isna(peps):
                if gene not in unknown_pep_warned:
                    unknown_pep_warned.add(gene)
                    _warn(line_no, "UNKNOWN_PEPTIDE_COUNT", gene,
                          f"{gene} is cited but its unique_peptides count is unknown; peptide-tier promotion cannot be verified.", line)
            elif int(peps) == 1 and not allow_untagged_single_peptide:
                if not any(tok in line_lower for tok in SINGLE_PEP_DISCLOSURE_TOKENS):
                    _err(line_no, "UNTAGGED_SINGLE_PEPTIDE_PROMOTION", gene,
                         f"Line {line_no} promotes {gene} (unique_peptides == 1) without an explicit single-peptide/exploratory disclosure.", line)

            # 3: overclaim lint
            if has_causal:
                weak = _gene_grade(ginfo) == "E3"
                g_fdr = pd.to_numeric(ginfo.get("empirical_null_fdr"), errors="coerce")
                if pd.notna(g_fdr) and float(g_fdr) > E3_FDR_THRESHOLD:
                    weak = True
                if ginfo.get("contrast") in weak_contrasts:
                    weak = True
                if weak:
                    _err(line_no, "OVERCLAIM_CAUSAL_LANGUAGE_ON_E3", gene,
                         f"Line {line_no} uses causal language for {gene}, whose evidence is E3 / empirical-null FDR > {E3_FDR_THRESHOLD}.", line)

        # 3b: causal verbs next to a weak contrast name
        if has_causal:
            for cname in weak_contrasts:
                if cname and cname in line:
                    _err(line_no, "OVERCLAIM_CAUSAL_LANGUAGE_ON_E3", cname,
                         f"Line {line_no} uses causal language for contrast '{cname}' (empirical-null FDR > {E3_FDR_THRESHOLD} / E3).", line)

        # 1A: signed gene-number reconciliation (nearest preceding gene within 60 chars)
        if mentions:
            for nm in _SIGNED_NUM.finditer(line):
                # ignore scientific notation exponents and ranges like "1.2-3.4" (handled by lookbehind)
                prev = [mm for mm in mentions if mm[1] <= nm.start() and nm.start() - mm[1] <= 60]
                if not prev:
                    continue
                gene = prev[-1][2]
                metrics = {k: v for k, v in (genes_db[gene].get("metrics") or {}).items()
                           if pd.notna(pd.to_numeric(v, errors="coerce"))}
                cited_vals = _to_float_signed(nm.group(1), nm.group(2))
                if not metrics:
                    _warn(line_no, "UNVERIFIABLE_NUMBER", gene,
                          f"Line {line_no} cites {nm.group(0)} for {gene}, but facts contain no metrics for {gene}.", line)
                    continue
                ok = any(abs(cv - float(mv)) <= tolerance + 1e-9 for cv in cited_vals for mv in metrics.values())
                if not ok:
                    _err(line_no, "GENE_NUMBER_MISMATCH", gene,
                         f"Line {line_no} cites {nm.group(0)} for {gene}; no canonical metric within ±{tolerance} "
                         f"(facts: {', '.join(f'{k}={float(v):+.2f}' for k, v in list(metrics.items())[:6])}).", line)

    # 4: citations
    text = "\n".join(lines)
    pmids = sorted(set(re.findall(r"PMID:?\s*(\d{5,9})", text, flags=re.IGNORECASE)))
    dois = sorted({d.rstrip(".,;)") for d in re.findall(r"10\.\d{4,9}/[-._;()/:A-Za-z0-9]+", text)})
    citation_checks: List[Dict[str, Any]] = []
    if verify_citations:
        for kind, ids in (("pmid", pmids), ("doi", dois)):
            for ident in ids:
                res = verify_citation_europepmc(kind, ident)
                citation_checks.append(res)
                if res["status"] == "NOT_FOUND":
                    _err(0, "CITATION_NOT_FOUND", ident, f"{kind.upper()} {ident} was not found in EuropePMC.", "")
                elif res["status"] == "UNVERIFIED" and not allow_offline:
                    _warn(0, "CITATION_UNVERIFIED", ident,
                          f"{kind.upper()} {ident} could not be verified (offline?): {res.get('reason', '')}")
    else:
        for kind, ids in (("pmid", pmids), ("doi", dois)):
            for ident in ids:
                citation_checks.append({"kind": kind, "id": ident, "status": "UNVERIFIED", "reason": "verification disabled"})
                if not allow_offline:
                    _warn(0, "CITATION_UNVERIFIED", ident, f"{kind.upper()} {ident} not verified (verification disabled).")

    status = "FAIL" if critical_errors else ("WARN" if warnings else "PASS")
    return {
        "file": source,
        "status": status,
        "n_critical_errors": len(critical_errors),
        "n_warnings": len(warnings),
        "critical_errors": critical_errors,
        "warnings": warnings,
        "extracted_pmids": pmids,
        "extracted_dois": dois,
        "citation_checks": citation_checks,
    }


def audit_markdown_against_facts(
    md_path: str | Path,
    facts: Dict[str, Any],
    allow_untagged_single_peptide: bool = False,
    tolerance: float = NUMBER_TOLERANCE,
    verify_citations: bool = True,
    allow_offline: bool = False,
) -> Dict[str, Any]:
    """Audit a Markdown/text report, SVG schematic, or PPTX deck against `facts`.

    Returns a dict with keys: file, status (PASS/WARN/FAIL), n_critical_errors,
    n_warnings, critical_errors, warnings, extracted_pmids, extracted_dois, citation_checks.
    """
    lines = extract_text_lines(md_path)
    return audit_text_lines_against_facts(
        lines, facts,
        allow_untagged_single_peptide=allow_untagged_single_peptide,
        tolerance=tolerance,
        verify_citations=verify_citations,
        allow_offline=allow_offline,
        source=str(md_path),
    )


audit_file_against_facts = audit_markdown_against_facts


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Stage 7B: Pre-Delivery Fact & Citation Gate")
    parser.add_argument("--report", "-r", required=True, nargs="+", help="Markdown / SVG / PPTX file(s) to audit")
    parser.add_argument("--facts-json", "-f", default=None, help="Canonical facts.json file")
    parser.add_argument("--csv", "-c", default=None, help="Or build facts directly from a canonical CSV file")
    parser.add_argument("--out", "-o", default=None, help="Optional JSON output report")
    parser.add_argument("--tolerance", type=float, default=NUMBER_TOLERANCE)
    parser.add_argument("--allow-offline", action="store_true", help="Do not WARN when citations cannot be verified offline")
    parser.add_argument("--no-verify-citations", action="store_true", help="Skip live EuropePMC lookups")
    args = parser.parse_args(argv)

    if args.facts_json:
        with open(args.facts_json, "r", encoding="utf-8") as f:
            facts = json.load(f)
    elif args.csv:
        facts = build_facts_dictionary_from_csv(args.csv)
    else:
        parser.error("Either --facts-json (-f) or --csv (-c) must be provided.")

    results = []
    for rp in args.report:
        res = audit_markdown_against_facts(
            rp, facts, tolerance=args.tolerance,
            verify_citations=not args.no_verify_citations, allow_offline=args.allow_offline,
        )
        results.append(res)
        print(f"[Stage 7B Fact Gate] {res['file']} -> {res['status']} "
              f"(Critical Errors: {res['n_critical_errors']}, Warnings: {res['n_warnings']})")
        for err in res["critical_errors"]:
            print(f"  [ERROR L{err['line']} {err['type']}] {err['gene']}: {err['message']}")
        for w in res["warnings"]:
            print(f"  [WARN  L{w['line']} {w['type']}] {w['gene']}: {w['message']}")

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(results if len(results) > 1 else results[0], f, indent=2, ensure_ascii=False)

    return 1 if any(r["status"] == "FAIL" for r in results) else 0


if __name__ == "__main__":
    sys.exit(main())
